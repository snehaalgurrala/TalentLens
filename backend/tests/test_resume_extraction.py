"""
Unit tests for app.services.resume_extraction.

No database, no HTTP, no AI — only file-system I/O via tmp_path.
Async tests run automatically because asyncio_mode = "auto" in pyproject.toml.
"""

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.resume_extraction import (
    ExtractionError,
    detect_file_type,
    extract_text_from_docx,
    extract_text_from_docx_bytes,
    extract_text_from_pdf,
    extract_text_from_pdf_bytes,
    extract_text_from_zip,
)

# ── In-memory fixture builders ────────────────────────────────────────────────


def _make_pdf() -> bytes:
    """Create a minimal valid (blank-page) PDF in memory via pypdf."""
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_docx(text: str | None = None) -> bytes:
    """Create a DOCX in memory; adds one paragraph when text is given."""
    doc = Document()
    if text:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_zip(files: dict[str, bytes]) -> bytes:
    """Build a ZIP in memory from {filename: content}."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# detect_file_type
# ─────────────────────────────────────────────────────────────────────────────


class TestDetectFileType:
    def test_pdf_lowercase(self):
        assert detect_file_type("resume.pdf") == "pdf"

    def test_pdf_uppercase(self):
        assert detect_file_type("RESUME.PDF") == "pdf"

    def test_docx(self):
        assert detect_file_type("resume.docx") == "docx"

    def test_docx_uppercase(self):
        assert detect_file_type("resume.DOCX") == "docx"

    def test_zip(self):
        assert detect_file_type("batch.zip") == "zip"

    def test_zip_uppercase(self):
        assert detect_file_type("BATCH.ZIP") == "zip"

    def test_unsupported_txt_raises(self):
        with pytest.raises(ValueError, match="Unsupported"):
            detect_file_type("notes.txt")

    def test_unsupported_png_raises(self):
        with pytest.raises(ValueError, match="Unsupported"):
            detect_file_type("photo.png")

    def test_no_extension_raises(self):
        with pytest.raises(ValueError, match="Unsupported"):
            detect_file_type("resume")

    def test_dot_only_raises(self):
        with pytest.raises(ValueError, match="Unsupported"):
            detect_file_type(".")


# ─────────────────────────────────────────────────────────────────────────────
# extract_text_from_pdf
# ─────────────────────────────────────────────────────────────────────────────


class TestExtractTextFromPdf:
    async def test_valid_pdf_returns_string(self, tmp_path):
        path = tmp_path / "test.pdf"
        path.write_bytes(_make_pdf())
        result = await extract_text_from_pdf(str(path))
        assert isinstance(result, str)

    async def test_blank_pdf_returns_empty_string(self, tmp_path):
        path = tmp_path / "blank.pdf"
        path.write_bytes(_make_pdf())
        result = await extract_text_from_pdf(str(path))
        # A blank page has no text content
        assert result == ""

    async def test_corrupted_pdf_raises(self, tmp_path):
        path = tmp_path / "corrupt.pdf"
        path.write_bytes(b"this is not a pdf at all")
        with pytest.raises(ExtractionError, match="Could not read PDF"):
            await extract_text_from_pdf(str(path))

    async def test_empty_file_raises(self, tmp_path):
        path = tmp_path / "empty.pdf"
        path.write_bytes(b"")
        with pytest.raises(ExtractionError):
            await extract_text_from_pdf(str(path))

    async def test_truncated_header_raises(self, tmp_path):
        path = tmp_path / "truncated.pdf"
        path.write_bytes(b"%PDF-1.4\n1 0 obj<<")  # header present but abruptly cut off
        with pytest.raises(ExtractionError):
            await extract_text_from_pdf(str(path))


# ─────────────────────────────────────────────────────────────────────────────
# extract_text_from_docx
# ─────────────────────────────────────────────────────────────────────────────


class TestExtractTextFromDocx:
    async def test_docx_with_text(self, tmp_path):
        path = tmp_path / "test.docx"
        path.write_bytes(_make_docx("Hello candidate"))
        result = await extract_text_from_docx(str(path))
        assert "Hello candidate" in result

    async def test_empty_docx_returns_empty_string(self, tmp_path):
        path = tmp_path / "empty.docx"
        path.write_bytes(_make_docx())  # no paragraphs added
        result = await extract_text_from_docx(str(path))
        assert result == ""

    async def test_multiple_paragraphs(self, tmp_path):
        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        buf = io.BytesIO()
        doc.save(buf)
        path = tmp_path / "multi.docx"
        path.write_bytes(buf.getvalue())
        result = await extract_text_from_docx(str(path))
        assert "First paragraph" in result
        assert "Second paragraph" in result

    async def test_corrupted_docx_raises(self, tmp_path):
        path = tmp_path / "corrupt.docx"
        path.write_bytes(b"not a docx file at all")
        with pytest.raises(ExtractionError, match="Could not read DOCX"):
            await extract_text_from_docx(str(path))

    async def test_pdf_masquerading_as_docx_raises(self, tmp_path):
        # DOCX is ZIP-based; a PDF is not, so python-docx raises on open
        path = tmp_path / "wrong_ext.docx"
        path.write_bytes(_make_pdf())
        with pytest.raises(ExtractionError, match="Could not read DOCX"):
            await extract_text_from_docx(str(path))

    async def test_empty_bytes_raises(self, tmp_path):
        path = tmp_path / "empty.docx"
        path.write_bytes(b"")
        with pytest.raises(ExtractionError):
            await extract_text_from_docx(str(path))


# ─────────────────────────────────────────────────────────────────────────────
# extract_text_from_zip
# ─────────────────────────────────────────────────────────────────────────────


class TestExtractTextFromZip:
    async def test_zip_with_pdf(self, tmp_path):
        path = tmp_path / "batch.zip"
        path.write_bytes(_make_zip({"cv.pdf": _make_pdf()}))
        results = await extract_text_from_zip(str(path))
        assert len(results) == 1
        assert results[0][0] == "cv.pdf"
        assert isinstance(results[0][1], str)

    async def test_zip_with_docx(self, tmp_path):
        path = tmp_path / "batch.zip"
        path.write_bytes(_make_zip({"alice.docx": _make_docx("Alice bio")}))
        results = await extract_text_from_zip(str(path))
        assert len(results) == 1
        assert results[0][0] == "alice.docx"
        assert "Alice bio" in results[0][1]

    async def test_zip_with_pdf_and_docx(self, tmp_path):
        path = tmp_path / "batch.zip"
        path.write_bytes(_make_zip({
            "cv.pdf": _make_pdf(),
            "bio.docx": _make_docx("Bio text"),
        }))
        results = await extract_text_from_zip(str(path))
        names = {r[0] for r in results}
        assert names == {"cv.pdf", "bio.docx"}

    async def test_zip_skips_unsupported_files(self, tmp_path):
        path = tmp_path / "mixed.zip"
        path.write_bytes(_make_zip({
            "cv.pdf": _make_pdf(),
            "notes.txt": b"ignore me",
            "photo.png": b"\x89PNG\r\n",
        }))
        results = await extract_text_from_zip(str(path))
        names = [r[0] for r in results]
        assert names == ["cv.pdf"]

    async def test_empty_zip_returns_empty_list(self, tmp_path):
        path = tmp_path / "empty.zip"
        path.write_bytes(_make_zip({}))
        results = await extract_text_from_zip(str(path))
        assert results == []

    async def test_zip_with_only_unsupported_returns_empty_list(self, tmp_path):
        path = tmp_path / "unsupported.zip"
        path.write_bytes(_make_zip({"data.csv": b"col1,col2\n1,2"}))
        results = await extract_text_from_zip(str(path))
        assert results == []

    async def test_nested_zip_raises(self, tmp_path):
        inner = _make_zip({"inner.pdf": _make_pdf()})
        path = tmp_path / "outer.zip"
        path.write_bytes(_make_zip({"inner.zip": inner}))
        with pytest.raises(ExtractionError, match="Nested ZIP"):
            await extract_text_from_zip(str(path))

    async def test_corrupted_zip_raises(self, tmp_path):
        path = tmp_path / "bad.zip"
        path.write_bytes(b"not a zip file at all")
        with pytest.raises(ExtractionError, match="not a valid ZIP"):
            await extract_text_from_zip(str(path))

    async def test_directory_entries_are_skipped_and_basename_used(self, tmp_path):
        # ZIP containing a directory entry + file nested inside it
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr(zipfile.ZipInfo("subdir/"), b"")  # directory entry
            zf.writestr("subdir/cv.pdf", _make_pdf())
        path = tmp_path / "dirs.zip"
        path.write_bytes(buf.getvalue())
        results = await extract_text_from_zip(str(path))
        assert len(results) == 1
        assert results[0][0] == "cv.pdf"  # basename only, no subdir/ prefix

    async def test_multiple_pdfs_all_extracted(self, tmp_path):
        path = tmp_path / "many.zip"
        path.write_bytes(_make_zip({
            "alice.pdf": _make_pdf(),
            "bob.pdf": _make_pdf(),
            "carol.pdf": _make_pdf(),
        }))
        results = await extract_text_from_zip(str(path))
        assert len(results) == 3
        names = {r[0] for r in results}
        assert names == {"alice.pdf", "bob.pdf", "carol.pdf"}


# ─────────────────────────────────────────────────────────────────────────────
# extract_text_from_pdf_bytes / extract_text_from_docx_bytes
# ─────────────────────────────────────────────────────────────────────────────


class TestExtractTextFromPdfBytes:
    async def test_valid_pdf_bytes_returns_string(self):
        result = await extract_text_from_pdf_bytes(_make_pdf(), "jd.pdf")
        assert isinstance(result, str)

    async def test_corrupted_pdf_bytes_raises(self):
        with pytest.raises(ExtractionError, match="Could not read PDF"):
            await extract_text_from_pdf_bytes(b"this is not a pdf at all", "jd.pdf")


class TestExtractTextFromDocxBytes:
    async def test_docx_bytes_with_text(self):
        result = await extract_text_from_docx_bytes(_make_docx("Senior Backend Engineer"), "jd.docx")
        assert "Senior Backend Engineer" in result

    async def test_corrupted_docx_bytes_raises(self):
        with pytest.raises(ExtractionError, match="Could not read DOCX"):
            await extract_text_from_docx_bytes(b"not a docx file at all", "jd.docx")
